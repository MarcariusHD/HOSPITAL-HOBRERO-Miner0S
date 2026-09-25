import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import os
import csv
import random
import threading
import unicodedata
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

from sklearn.model_selection import train_test_split
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    precision_score,
    recall_score,
    f1_score,
    confusion_matrix
)


# ============================================================
# CONFIGURACIÓN
# ============================================================

archivo_actual = None

# Para el modelo NO necesitamos cargar 1 millón de filas.
# Una muestra de 80.000 suele ser más que suficiente para
# una primera regresión logística y acelera muchísimo el proceso.
MAX_FILAS_MODELO = 80000

# Tamaño de bloque para CSV.
CHUNK_CSV = 100000

# Semilla reproducible.
SEMILLA = 42

# Colores de interfaz
COLOR_FONDO = "#F4F7FB"
COLOR_PANEL = "#FFFFFF"
COLOR_PRIMARIO = "#2563EB"
COLOR_TEXTO = "#172033"
COLOR_SECUNDARIO = "#64748B"
COLOR_EXITO = "#15803D"
COLOR_BORDE = "#DCE3ED"
COLOR_SUAVE = "#EEF4FF"


# ============================================================
# UTILIDADES
# ============================================================

def normalizar_texto(texto):
    texto = str(texto).strip().lower()
    texto = unicodedata.normalize("NFKD", texto)

    return "".join(
        c for c in texto
        if not unicodedata.combining(c)
    )


def buscar_columna(columnas, opciones):
    mapa = {
        normalizar_texto(columna): columna
        for columna in columnas
    }

    for opcion in opciones:
        clave = normalizar_texto(opcion)

        if clave in mapa:
            return mapa[clave]

    return None


def detectar_columnas(columnas):
    """
    Detecta automáticamente los encabezados que necesitamos.
    """

    col_triaje = buscar_columna(
        columnas,
        [
            "Nivel de triaje",
            "Nivel de triage",
            "Nivel Triaje",
            "Triaje",
            "Triage"
        ]
    )

    col_edad = buscar_columna(
        columnas,
        ["Edad"]
    )

    col_sexo = buscar_columna(
        columnas,
        ["Sexo", "Genero", "Género"]
    )

    col_hora = buscar_columna(
        columnas,
        ["Hora"]
    )

    col_dia = buscar_columna(
        columnas,
        ["Día", "Dia"]
    )

    if col_triaje is None:
        raise ValueError(
            "No se encontró la columna de triaje.\n\n"
            "Se buscó: 'Nivel de triaje', 'Triaje', "
            "'Nivel de triage' o 'Triage'."
        )

    return {
        "triaje": col_triaje,
        "edad": col_edad,
        "sexo": col_sexo,
        "hora": col_hora,
        "dia": col_dia
    }


def triaje_valido(valor):
    try:
        numero = int(float(valor))

        if numero in (1, 2, 3, 4, 5):
            return numero

    except (TypeError, ValueError):
        pass

    return None


def edad_valida(valor):
    try:
        numero = float(valor)

        if 0 <= numero <= 120:
            return numero

    except (TypeError, ValueError):
        pass

    return None


def hora_a_numero(valor):
    """
    Convierte una hora a un número de 0 a 23.

    Acepta:
    - 14
    - "14:35"
    - objetos datetime/time de Excel
    """

    if valor is None:
        return None

    if hasattr(valor, "hour"):
        return int(valor.hour)

    try:
        numero = float(valor)

        # Si Excel guarda la hora como fracción del día.
        if 0 <= numero < 1:
            return int(numero * 24)

        if 0 <= numero <= 23:
            return int(numero)

    except (TypeError, ValueError):
        pass

    texto = str(valor).strip()

    if not texto:
        return None

    try:
        hora = int(texto.split(":")[0])

        if 0 <= hora <= 23:
            return hora

    except (TypeError, ValueError):
        pass

    return None


def limpiar_categoria(valor):
    if valor is None:
        return None

    texto = str(valor).strip()

    if not texto or texto.lower() in {
        "nan", "none", "null"
    }:
        return None

    return texto


def agregar_muestra_reservorio(
    muestra,
    registro,
    total_registros_validos
):
    """
    Reservoir Sampling:
    permite obtener una muestra aleatoria de un archivo enorme
    SIN guardar todo el archivo en memoria.

    Si todavía hay espacio, guarda el registro.
    Si ya se llenó la muestra, decide aleatoriamente si reemplaza
    alguno de los registros guardados.
    """

    if len(muestra) < MAX_FILAS_MODELO:
        muestra.append(registro)
        return

    posicion = random.randint(
        0,
        total_registros_validos - 1
    )

    if posicion < MAX_FILAS_MODELO:
        muestra[posicion] = registro


def actualizar_estado(texto):
    ventana.after(
        0,
        lambda: lbl_estado.config(text=texto)
    )


def actualizar_progreso(valor):
    ventana.after(
        0,
        lambda: progreso.configure(value=valor)
    )


# ============================================================
# SELECCIÓN DE ARCHIVO
# ============================================================

def seleccionar_archivo():
    global archivo_actual

    ruta = filedialog.askopenfilename(
        title="Seleccionar dataset",
        filetypes=[
            ("Archivos de datos", "*.xlsx *.csv"),
            ("Excel", "*.xlsx"),
            ("CSV", "*.csv"),
            ("Todos", "*.*")
        ]
    )

    if not ruta:
        return

    archivo_actual = ruta

    nombre = os.path.basename(ruta)
    tam_mb = os.path.getsize(ruta) / (1024 * 1024)

    lbl_archivo_nombre.config(
        text=nombre
    )

    lbl_archivo_info.config(
        text=f"{tam_mb:.1f} MB · Listo para analizar"
    )

    lbl_estado.config(
        text="Archivo seleccionado. Presiona «Analizar dataset»."
    )

    actualizar_progreso(0)
    limpiar_resultados()


# ============================================================
# PROCESAMIENTO CSV
# ============================================================

def procesar_csv(ruta):
    """
    Lee CSV por bloques.

    Ventaja:
    no carga todo el archivo en RAM.
    """

    encabezado = pd.read_csv(
        ruta,
        nrows=0
    )

    columnas = detectar_columnas(
        encabezado.columns.tolist()
    )

    columnas_usar = [
        valor
        for valor in columnas.values()
        if valor is not None
    ]

    columnas_usar = list(
        dict.fromkeys(columnas_usar)
    )

    conteos = Counter()
    muestra = []

    total_validos = 0
    total_modelo_validos = 0
    filas_leidas = 0

    tamaño_archivo = max(
        os.path.getsize(ruta),
        1
    )

    # Para CSV podemos usar chunks, que es muy rápido.
    for chunk in pd.read_csv(
        ruta,
        usecols=columnas_usar,
        chunksize=CHUNK_CSV,
        low_memory=False
    ):
        filas_leidas += len(chunk)

        # ----------------------------
        # CONTEO DE TRIAJE VECTORIAL
        # ----------------------------
        triaje = pd.to_numeric(
            chunk[columnas["triaje"]],
            errors="coerce"
        )

        validos = triaje.isin(
            [1, 2, 3, 4, 5]
        )

        triaje_validos = (
            triaje.loc[validos]
            .astype("int8")
        )

        conteo_chunk = (
            triaje_validos
            .value_counts()
            .to_dict()
        )

        for nivel, cantidad in conteo_chunk.items():
            conteos[int(nivel)] += int(cantidad)

        total_validos += len(
            triaje_validos
        )

        # ----------------------------
        # CREAR MUESTRA PARA ML
        # ----------------------------
        sub = chunk.loc[validos].copy()

        sub["_triaje"] = (
            triaje.loc[validos]
            .astype("int8")
        )

        # Convertimos solo las variables disponibles.
        if columnas["edad"]:
            sub["_edad"] = pd.to_numeric(
                sub[columnas["edad"]],
                errors="coerce"
            )
        else:
            sub["_edad"] = pd.NA

        if columnas["hora"]:
            # Conversión vectorial primero.
            hora_num = pd.to_numeric(
                sub[columnas["hora"]],
                errors="coerce"
            )

            # Solo para los casos que no se pudieron convertir,
            # aplicamos la función más flexible.
            faltan = hora_num.isna()

            if faltan.any():
                hora_num.loc[faltan] = (
                    sub.loc[
                        faltan,
                        columnas["hora"]
                    ]
                    .map(hora_a_numero)
                )

            sub["_hora"] = hora_num
        else:
            sub["_hora"] = pd.NA

        if columnas["sexo"]:
            sub["_sexo"] = (
                sub[columnas["sexo"]]
                .astype("string")
                .str.strip()
            )
        else:
            sub["_sexo"] = pd.NA

        if columnas["dia"]:
            sub["_dia"] = (
                sub[columnas["dia"]]
                .astype("string")
                .str.strip()
            )
        else:
            sub["_dia"] = pd.NA

        # Seleccionamos solo las columnas internas.
        sub = sub[
            [
                "_triaje",
                "_edad",
                "_hora",
                "_sexo",
                "_dia"
            ]
        ]

        # No obligamos a que existan TODAS.
        # Luego elegiremos solo variables disponibles.
        registros = sub.to_dict(
            orient="records"
        )

        for registro in registros:
            total_modelo_validos += 1

            agregar_muestra_reservorio(
                muestra,
                registro,
                total_modelo_validos
            )

        # Barra aproximada.
        # En CSV podemos usar el progreso por filas/chunks.
        avance = min(
            70,
            5 + int(
                60 * (
                    filas_leidas /
                    max(filas_leidas + CHUNK_CSV, 1)
                )
            )
        )

        actualizar_progreso(avance)

        actualizar_estado(
            f"Leyendo CSV… {filas_leidas:,} filas procesadas"
        )

    return (
        conteos,
        total_validos,
        muestra,
        columnas
    )


# ============================================================
# PROCESAMIENTO EXCEL OPTIMIZADO
# ============================================================

def procesar_excel(ruta):
    """
    Lee el Excel en modo READ-ONLY.

    Esto es MUCHO más eficiente en memoria que pd.read_excel()
    para archivos grandes porque no crea un DataFrame de
    1 millón de filas.
    """

    actualizar_estado(
        "Abriendo Excel en modo optimizado…"
    )

    libro = load_workbook(
        ruta,
        read_only=True,
        data_only=True
    )

    hoja = libro.active

    iterador = hoja.iter_rows(
        values_only=True
    )

    try:
        encabezados = next(iterador)
    except StopIteration:
        libro.close()
        raise ValueError(
            "El archivo Excel está vacío."
        )

    encabezados = [
        str(valor).strip()
        if valor is not None
        else ""
        for valor in encabezados
    ]

    columnas = detectar_columnas(
        encabezados
    )

    # Índice de cada columna dentro de la fila.
    indices = {}

    for clave, nombre in columnas.items():
        if nombre is not None:
            indices[clave] = (
                encabezados.index(nombre)
            )
        else:
            indices[clave] = None

    conteos = Counter()
    muestra = []

    total_validos = 0
    total_modelo_validos = 0

    # max_row de openpyxl nos permite mostrar progreso aproximado.
    total_filas_estimadas = max(
        hoja.max_row - 1,
        1
    )

    for numero_fila, fila in enumerate(
        iterador,
        start=1
    ):
        indice_triaje = indices["triaje"]

        if indice_triaje >= len(fila):
            continue

        nivel = triaje_valido(
            fila[indice_triaje]
        )

        if nivel is None:
            continue

        conteos[nivel] += 1
        total_validos += 1

        registro = {
            "_triaje": nivel,
            "_edad": None,
            "_hora": None,
            "_sexo": None,
            "_dia": None
        }

        if indices["edad"] is not None:
            idx = indices["edad"]

            if idx < len(fila):
                registro["_edad"] = edad_valida(
                    fila[idx]
                )

        if indices["hora"] is not None:
            idx = indices["hora"]

            if idx < len(fila):
                registro["_hora"] = hora_a_numero(
                    fila[idx]
                )

        if indices["sexo"] is not None:
            idx = indices["sexo"]

            if idx < len(fila):
                registro["_sexo"] = limpiar_categoria(
                    fila[idx]
                )

        if indices["dia"] is not None:
            idx = indices["dia"]

            if idx < len(fila):
                registro["_dia"] = limpiar_categoria(
                    fila[idx]
                )

        total_modelo_validos += 1

        agregar_muestra_reservorio(
            muestra,
            registro,
            total_modelo_validos
        )

        # No actualizar la interfaz por cada fila:
        # hacerlo cada 10.000 evita ralentizar el proceso.
        if numero_fila % 10000 == 0:
            porcentaje = (
                numero_fila /
                total_filas_estimadas
            )

            avance = min(
                70,
                5 + int(porcentaje * 65)
            )

            actualizar_progreso(
                avance
            )

            actualizar_estado(
                f"Leyendo Excel… {numero_fila:,} filas procesadas"
            )

    libro.close()

    return (
        conteos,
        total_validos,
        muestra,
        columnas
    )


# ============================================================
# PREPARACIÓN Y REGRESIÓN LOGÍSTICA
# ============================================================

def entrenar_modelo(muestra, columnas):
    actualizar_estado(
        "Preparando datos para Machine Learning…"
    )
    actualizar_progreso(75)

    if not muestra:
        raise ValueError(
            "No hay registros válidos para entrenar el modelo."
        )

    df = pd.DataFrame(
        muestra
    )

    # Target:
    # 0 = Triaje 1, 2, 3
    # 1 = Triaje 4, 5
    df["objetivo"] = (
        df["_triaje"]
        .isin([4, 5])
        .astype("int8")
    )

    variables = []

    # Añadimos una variable solo si realmente tiene datos suficientes.
    if columnas["edad"] is not None:
        df["_edad"] = pd.to_numeric(
            df["_edad"],
            errors="coerce"
        )

        if df["_edad"].notna().sum() >= 50:
            variables.append("_edad")

    if columnas["hora"] is not None:
        df["_hora"] = pd.to_numeric(
            df["_hora"],
            errors="coerce"
        )

        if df["_hora"].notna().sum() >= 50:
            variables.append("_hora")

    if columnas["sexo"] is not None:
        if df["_sexo"].notna().sum() >= 50:
            variables.append("_sexo")

    if columnas["dia"] is not None:
        if df["_dia"].notna().sum() >= 50:
            variables.append("_dia")

    if not variables:
        raise ValueError(
            "No se encontraron suficientes datos en Edad, Hora, "
            "Sexo o Día para entrenar la regresión logística."
        )

    trabajo = df[
        variables + ["objetivo"]
    ].copy()

    # Para no tirar una fila porque falta UNA variable categórica,
    # usamos imputación sencilla:
    # - numéricas: mediana
    # - categóricas: "Desconocido"
    for columna in variables:
        if columna in ("_edad", "_hora"):
            mediana = trabajo[
                columna
            ].median()

            trabajo[columna] = (
                trabajo[columna]
                .fillna(mediana)
            )
        else:
            trabajo[columna] = (
                trabajo[columna]
                .fillna("Desconocido")
                .astype(str)
                .str.strip()
            )

    columnas_categoricas = [
        columna
        for columna in variables
        if columna in ("_sexo", "_dia")
    ]

    X = trabajo[
        variables
    ].copy()

    if columnas_categoricas:
        X = pd.get_dummies(
            X,
            columns=columnas_categoricas,
            drop_first=True,
            dtype="int8"
        )

    # Convertir a float32 reduce memoria.
    X = X.astype("float32")

    y = trabajo[
        "objetivo"
    ].astype("int8")

    if y.nunique() < 2:
        raise ValueError(
            "Solo se encontró una clase de triaje.\n"
            "Se necesitan registros tanto de 1-3 como de 4-5."
        )

    actualizar_estado(
        "Dividiendo datos: 80% entrenamiento / 20% prueba…"
    )
    actualizar_progreso(82)

    X_train, X_test, y_train, y_test = train_test_split(
        X,
        y,
        test_size=0.20,
        random_state=SEMILLA,
        stratify=y
    )

    actualizar_estado(
        "Entrenando regresión logística…"
    )
    actualizar_progreso(88)

    # liblinear es muy eficiente para clasificación binaria
    # cuando tenemos pocas variables.
    modelo = LogisticRegression(
        solver="liblinear",
        class_weight="balanced",
        max_iter=300,
        random_state=SEMILLA
    )

    modelo.fit(
        X_train,
        y_train
    )

    actualizar_estado(
        "Evaluando el modelo…"
    )
    actualizar_progreso(94)

    pred = modelo.predict(
        X_test
    )

    accuracy = accuracy_score(
        y_test,
        pred
    )

    precision = precision_score(
        y_test,
        pred,
        zero_division=0
    )

    recall = recall_score(
        y_test,
        pred,
        zero_division=0
    )

    f1 = f1_score(
        y_test,
        pred,
        zero_division=0
    )

    matriz = confusion_matrix(
        y_test,
        pred,
        labels=[0, 1]
    )

    tn, fp, fn, tp = matriz.ravel()

    nombres_legibles = []

    for variable in variables:
        if variable == "_edad":
            nombres_legibles.append("Edad")
        elif variable == "_hora":
            nombres_legibles.append("Hora")
        elif variable == "_sexo":
            nombres_legibles.append("Sexo")
        elif variable == "_dia":
            nombres_legibles.append("Día")

    return {
        "accuracy": accuracy,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "tn": int(tn),
        "fp": int(fp),
        "fn": int(fn),
        "tp": int(tp),
        "train": len(X_train),
        "test": len(X_test),
        "muestra": len(df),
        "variables": nombres_legibles
    }


# ============================================================
# PROCESO COMPLETO
# ============================================================

def procesar_archivo():
    random.seed(SEMILLA)

    try:
        ruta = archivo_actual

        actualizar_progreso(3)

        if ruta.lower().endswith(".csv"):
            (
                conteos,
                total_validos,
                muestra,
                columnas
            ) = procesar_csv(ruta)

        elif ruta.lower().endswith(".xlsx"):
            (
                conteos,
                total_validos,
                muestra,
                columnas
            ) = procesar_excel(ruta)

        else:
            raise ValueError(
                "Formato no compatible. Usa .csv o .xlsx."
            )

        actualizar_estado(
            "Calculando resultados descriptivos…"
        )
        actualizar_progreso(72)

        resultados = []

        for nivel in range(1, 6):
            cantidad = int(
                conteos.get(nivel, 0)
            )

            porcentaje = (
                cantidad / total_validos * 100
                if total_validos
                else 0
            )

            resultados.append(
                (
                    nivel,
                    cantidad,
                    porcentaje
                )
            )

        resultados_modelo = entrenar_modelo(
            muestra,
            columnas
        )

        actualizar_progreso(100)

        ventana.after(
            0,
            mostrar_resultados,
            resultados,
            total_validos,
            resultados_modelo
        )

    except Exception as error:
        ventana.after(
            0,
            mostrar_error,
            str(error)
        )


# ============================================================
# INTERFAZ - RESULTADOS
# ============================================================

def limpiar_resultados():
    for item in tabla.get_children():
        tabla.delete(item)

    valor_total.config(text="—")
    valor_45.config(text="—")
    valor_accuracy.config(text="—")
    valor_precision.config(text="—")
    valor_recall.config(text="—")
    valor_f1.config(text="—")

    lbl_train_test.config(
        text="Ejecuta el análisis para ver los resultados."
    )

    lbl_variables.config(
        text="Variables del modelo: —"
    )

    for item in tabla_matriz.get_children():
        tabla_matriz.delete(item)


def mostrar_resultados(
    resultados,
    total,
    modelo
):
    btn_analizar.config(
        state="normal"
    )
    btn_archivo.config(
        state="normal"
    )

    porcentaje_45 = 0

    for item in tabla.get_children():
        tabla.delete(item)

    for nivel, cantidad, porcentaje in resultados:
        tabla.insert(
            "",
            "end",
            values=(
                f"Nivel {nivel}",
                f"{cantidad:,}",
                f"{porcentaje:.2f}%"
            )
        )

        if nivel in (4, 5):
            porcentaje_45 += porcentaje

    valor_total.config(
        text=f"{total:,}"
    )

    valor_45.config(
        text=f"{porcentaje_45:.2f}%"
    )

    valor_accuracy.config(
        text=f"{modelo['accuracy'] * 100:.2f}%"
    )

    valor_precision.config(
        text=f"{modelo['precision'] * 100:.2f}%"
    )

    valor_recall.config(
        text=f"{modelo['recall'] * 100:.2f}%"
    )

    valor_f1.config(
        text=f"{modelo['f1'] * 100:.2f}%"
    )

    lbl_train_test.config(
        text=(
            f"Muestra ML: {modelo['muestra']:,} registros  ·  "
            f"Train: {modelo['train']:,}  ·  "
            f"Test: {modelo['test']:,}"
        )
    )

    lbl_variables.config(
        text=(
            "Variables utilizadas: "
            + ", ".join(modelo["variables"])
        )
    )

    for item in tabla_matriz.get_children():
        tabla_matriz.delete(item)

    tabla_matriz.insert(
        "",
        "end",
        values=(
            "Real 0",
            modelo["tn"],
            modelo["fp"]
        )
    )

    tabla_matriz.insert(
        "",
        "end",
        values=(
            "Real 1",
            modelo["fn"],
            modelo["tp"]
        )
    )

    lbl_estado.config(
        text="✓ Análisis completado correctamente."
    )

    estado_dot.config(
        text="●",
        foreground=COLOR_EXITO
    )


def mostrar_error(error):
    btn_analizar.config(
        state="normal"
    )
    btn_archivo.config(
        state="normal"
    )

    progreso.configure(
        value=0
    )

    lbl_estado.config(
        text="Ocurrió un error durante el análisis."
    )

    estado_dot.config(
        text="●",
        foreground="#DC2626"
    )

    messagebox.showerror(
        "Error",
        error
    )


# ============================================================
# EVENTO ANALIZAR
# ============================================================

def analizar():
    if archivo_actual is None:
        messagebox.showwarning(
            "Selecciona un archivo",
            "Primero selecciona un archivo Excel o CSV."
        )
        return

    btn_analizar.config(
        state="disabled"
    )
    btn_archivo.config(
        state="disabled"
    )

    limpiar_resultados()

    estado_dot.config(
        text="●",
        foreground="#F59E0B"
    )

    lbl_estado.config(
        text="Preparando análisis…"
    )

    progreso.configure(
        value=1
    )

    hilo = threading.Thread(
        target=procesar_archivo,
        daemon=True
    )

    hilo.start()


# ============================================================
# INTERFAZ - HELPERS
# ============================================================

def crear_tarjeta(
    padre,
    titulo,
    valor_inicial="—",
    descripcion=""
):
    frame = tk.Frame(
        padre,
        bg=COLOR_PANEL,
        highlightbackground=COLOR_BORDE,
        highlightthickness=1
    )

    tk.Label(
        frame,
        text=titulo,
        bg=COLOR_PANEL,
        fg=COLOR_SECUNDARIO,
        font=("Segoe UI", 9, "bold")
    ).pack(
        anchor="w",
        padx=18,
        pady=(15, 4)
    )

    valor = tk.Label(
        frame,
        text=valor_inicial,
        bg=COLOR_PANEL,
        fg=COLOR_TEXTO,
        font=("Segoe UI", 21, "bold")
    )

    valor.pack(
        anchor="w",
        padx=18
    )

    if descripcion:
        tk.Label(
            frame,
            text=descripcion,
            bg=COLOR_PANEL,
            fg=COLOR_SECUNDARIO,
            font=("Segoe UI", 8),
            wraplength=165,
            justify="left"
        ).pack(
            anchor="w",
            padx=18,
            pady=(5, 14)
        )
    else:
        tk.Frame(
            frame,
            bg=COLOR_PANEL,
            height=12
        ).pack()

    return frame, valor


def rueda_mouse(evento):
    canvas.yview_scroll(
        int(-1 * (evento.delta / 120)),
        "units"
    )


# ============================================================
# INTERFAZ PRINCIPAL
# ============================================================

ventana = tk.Tk()

ventana.title(
    "Triaje Analytics · Minería de Datos"
)

ventana.geometry(
    "1000x780"
)

ventana.minsize(
    900,
    680
)

ventana.configure(
    bg=COLOR_FONDO
)

# ----------------------------
# ESTILOS TTK
# ----------------------------

style = ttk.Style()

try:
    style.theme_use("clam")
except tk.TclError:
    pass

style.configure(
    "Treeview",
    font=("Segoe UI", 10),
    rowheight=34,
    background=COLOR_PANEL,
    fieldbackground=COLOR_PANEL,
    foreground=COLOR_TEXTO,
    borderwidth=0
)

style.configure(
    "Treeview.Heading",
    font=("Segoe UI", 10, "bold"),
    background=COLOR_SUAVE,
    foreground=COLOR_TEXTO,
    relief="flat"
)

style.map(
    "Treeview",
    background=[
        ("selected", "#DBEAFE")
    ],
    foreground=[
        ("selected", COLOR_TEXTO)
    ]
)

style.configure(
    "Modern.Horizontal.TProgressbar",
    troughcolor="#E7EDF5",
    background=COLOR_PRIMARIO,
    bordercolor="#E7EDF5",
    lightcolor=COLOR_PRIMARIO,
    darkcolor=COLOR_PRIMARIO
)

# ----------------------------
# SCROLL PRINCIPAL
# ----------------------------

canvas = tk.Canvas(
    ventana,
    bg=COLOR_FONDO,
    highlightthickness=0
)

scroll = ttk.Scrollbar(
    ventana,
    orient="vertical",
    command=canvas.yview
)

contenido = tk.Frame(
    canvas,
    bg=COLOR_FONDO
)

contenido_id = canvas.create_window(
    (0, 0),
    window=contenido,
    anchor="nw"
)

canvas.configure(
    yscrollcommand=scroll.set
)

canvas.pack(
    side="left",
    fill="both",
    expand=True
)

scroll.pack(
    side="right",
    fill="y"
)

def ajustar_contenido(evento):
    canvas.itemconfigure(
        contenido_id,
        width=evento.width
    )

canvas.bind(
    "<Configure>",
    ajustar_contenido
)

contenido.bind(
    "<Configure>",
    lambda e: canvas.configure(
        scrollregion=canvas.bbox("all")
    )
)

canvas.bind_all(
    "<MouseWheel>",
    rueda_mouse
)

# ----------------------------
# CABECERA
# ----------------------------

header = tk.Frame(
    contenido,
    bg=COLOR_PRIMARIO,
    height=145
)

header.pack(
    fill="x"
)

header.pack_propagate(
    False
)

header_inner = tk.Frame(
    header,
    bg=COLOR_PRIMARIO
)

header_inner.pack(
    fill="both",
    expand=True,
    padx=42,
    pady=25
)

tk.Label(
    header_inner,
    text="Triaje Analytics",
    bg=COLOR_PRIMARIO,
    fg="white",
    font=("Segoe UI", 25, "bold")
).pack(
    anchor="w"
)

tk.Label(
    header_inner,
    text=(
        "Análisis descriptivo + Regresión Logística "
        "para datasets de urgencias"
    ),
    bg=COLOR_PRIMARIO,
    fg="#DBEAFE",
    font=("Segoe UI", 11)
).pack(
    anchor="w",
    pady=(4, 0)
)

# ----------------------------
# ZONA DE TRABAJO
# ----------------------------

main = tk.Frame(
    contenido,
    bg=COLOR_FONDO
)

main.pack(
    fill="both",
    expand=True,
    padx=42,
    pady=28
)

# ----------------------------
# TARJETA ARCHIVO
# ----------------------------

archivo_card = tk.Frame(
    main,
    bg=COLOR_PANEL,
    highlightbackground=COLOR_BORDE,
    highlightthickness=1
)

archivo_card.pack(
    fill="x",
    pady=(0, 18)
)

archivo_left = tk.Frame(
    archivo_card,
    bg=COLOR_PANEL
)

archivo_left.pack(
    side="left",
    fill="both",
    expand=True,
    padx=22,
    pady=18
)

tk.Label(
    archivo_left,
    text="Dataset",
    bg=COLOR_PANEL,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 12, "bold")
).pack(
    anchor="w"
)

lbl_archivo_nombre = tk.Label(
    archivo_left,
    text="Ningún archivo seleccionado",
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 10)
)

lbl_archivo_nombre.pack(
    anchor="w",
    pady=(5, 1)
)

lbl_archivo_info = tk.Label(
    archivo_left,
    text="Compatible con .xlsx y .csv",
    bg=COLOR_PANEL,
    fg="#94A3B8",
    font=("Segoe UI", 9)
)

lbl_archivo_info.pack(
    anchor="w"
)

acciones = tk.Frame(
    archivo_card,
    bg=COLOR_PANEL
)

acciones.pack(
    side="right",
    padx=22,
    pady=18
)

btn_archivo = tk.Button(
    acciones,
    text="Elegir archivo",
    command=seleccionar_archivo,
    font=("Segoe UI", 10, "bold"),
    bg="#E8EEF9",
    fg=COLOR_TEXTO,
    activebackground="#DCE7F7",
    activeforeground=COLOR_TEXTO,
    relief="flat",
    bd=0,
    padx=18,
    pady=10,
    cursor="hand2"
)

btn_archivo.pack(
    side="left",
    padx=(0, 10)
)

btn_analizar = tk.Button(
    acciones,
    text="Analizar dataset",
    command=analizar,
    font=("Segoe UI", 10, "bold"),
    bg=COLOR_PRIMARIO,
    fg="white",
    activebackground="#1D4ED8",
    activeforeground="white",
    relief="flat",
    bd=0,
    padx=20,
    pady=10,
    cursor="hand2"
)

btn_analizar.pack(
    side="left"
)

# ----------------------------
# ESTADO / PROGRESO
# ----------------------------

estado_card = tk.Frame(
    main,
    bg=COLOR_PANEL,
    highlightbackground=COLOR_BORDE,
    highlightthickness=1
)

estado_card.pack(
    fill="x",
    pady=(0, 18)
)

estado_top = tk.Frame(
    estado_card,
    bg=COLOR_PANEL
)

estado_top.pack(
    fill="x",
    padx=20,
    pady=(14, 8)
)

estado_dot = tk.Label(
    estado_top,
    text="●",
    bg=COLOR_PANEL,
    fg="#94A3B8",
    font=("Segoe UI", 10)
)

estado_dot.pack(
    side="left"
)

lbl_estado = tk.Label(
    estado_top,
    text="Esperando un archivo.",
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9)
)

lbl_estado.pack(
    side="left",
    padx=(7, 0)
)

progreso = ttk.Progressbar(
    estado_card,
    mode="determinate",
    maximum=100,
    style="Modern.Horizontal.TProgressbar"
)

progreso.pack(
    fill="x",
    padx=20,
    pady=(0, 16)
)

# ----------------------------
# MÉTRICAS PRINCIPALES
# ----------------------------

metricas_frame = tk.Frame(
    main,
    bg=COLOR_FONDO
)

metricas_frame.pack(
    fill="x",
    pady=(0, 18)
)

for i in range(4):
    metricas_frame.grid_columnconfigure(
        i,
        weight=1,
        uniform="metric"
    )

card_total, valor_total = crear_tarjeta(
    metricas_frame,
    "REGISTROS VÁLIDOS",
    descripcion="Filas con nivel de triaje entre 1 y 5."
)

card_total.grid(
    row=0,
    column=0,
    sticky="nsew",
    padx=(0, 7)
)

card_45, valor_45 = crear_tarjeta(
    metricas_frame,
    "TRIAJE 4 + 5",
    descripcion="Proporción conjunta de los niveles 4 y 5."
)

card_45.grid(
    row=0,
    column=1,
    sticky="nsew",
    padx=7
)

card_acc, valor_accuracy = crear_tarjeta(
    metricas_frame,
    "ACCURACY",
    descripcion="Porcentaje total de predicciones correctas."
)

card_acc.grid(
    row=0,
    column=2,
    sticky="nsew",
    padx=7
)

card_f1, valor_f1 = crear_tarjeta(
    metricas_frame,
    "F1-SCORE",
    descripcion="Equilibrio entre Precision y Recall."
)

card_f1.grid(
    row=0,
    column=3,
    sticky="nsew",
    padx=(7, 0)
)

# ----------------------------
# NOTEBOOK
# ----------------------------

notebook = ttk.Notebook(
    main
)

notebook.pack(
    fill="both",
    expand=True
)

tab_triaje = tk.Frame(
    notebook,
    bg=COLOR_PANEL
)

tab_modelo = tk.Frame(
    notebook,
    bg=COLOR_PANEL
)

notebook.add(
    tab_triaje,
    text="  Distribución de triaje  "
)

notebook.add(
    tab_modelo,
    text="  Modelo predictivo  "
)

# ----------------------------
# TAB TRIAJE
# ----------------------------

triaje_inner = tk.Frame(
    tab_triaje,
    bg=COLOR_PANEL
)

triaje_inner.pack(
    fill="both",
    expand=True,
    padx=24,
    pady=22
)

tk.Label(
    triaje_inner,
    text="Distribución por nivel",
    bg=COLOR_PANEL,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 14, "bold")
).pack(
    anchor="w"
)

tk.Label(
    triaje_inner,
    text=(
        "Conteo y porcentaje de registros válidos "
        "para cada nivel de triaje."
    ),
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9)
).pack(
    anchor="w",
    pady=(3, 14)
)

tabla = ttk.Treeview(
    triaje_inner,
    columns=("nivel", "cantidad", "porcentaje"),
    show="headings",
    height=6
)

tabla.heading(
    "nivel",
    text="Nivel"
)

tabla.heading(
    "cantidad",
    text="Cantidad"
)

tabla.heading(
    "porcentaje",
    text="Porcentaje"
)

tabla.column(
    "nivel",
    anchor="center",
    width=180
)

tabla.column(
    "cantidad",
    anchor="center",
    width=220
)

tabla.column(
    "porcentaje",
    anchor="center",
    width=220
)

tabla.pack(
    fill="x"
)

# ----------------------------
# TAB MODELO
# ----------------------------

modelo_inner = tk.Frame(
    tab_modelo,
    bg=COLOR_PANEL
)

modelo_inner.pack(
    fill="both",
    expand=True,
    padx=24,
    pady=22
)

tk.Label(
    modelo_inner,
    text="Regresión Logística",
    bg=COLOR_PANEL,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 14, "bold")
).pack(
    anchor="w"
)

tk.Label(
    modelo_inner,
    text=(
        "Clase 0 = Triaje 1, 2 o 3  ·  "
        "Clase 1 = Triaje 4 o 5"
    ),
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9)
).pack(
    anchor="w",
    pady=(3, 14)
)

modelo_metricas = tk.Frame(
    modelo_inner,
    bg=COLOR_PANEL
)

modelo_metricas.pack(
    fill="x",
    pady=(0, 15)
)

for i in range(2):
    modelo_metricas.grid_columnconfigure(
        i,
        weight=1
    )

precision_frame = tk.Frame(
    modelo_metricas,
    bg=COLOR_SUAVE
)

precision_frame.grid(
    row=0,
    column=0,
    sticky="ew",
    padx=(0, 6)
)

tk.Label(
    precision_frame,
    text="Precision",
    bg=COLOR_SUAVE,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9, "bold")
).pack(
    anchor="w",
    padx=16,
    pady=(12, 2)
)

valor_precision = tk.Label(
    precision_frame,
    text="—",
    bg=COLOR_SUAVE,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 17, "bold")
)

valor_precision.pack(
    anchor="w",
    padx=16,
    pady=(0, 12)
)

recall_frame = tk.Frame(
    modelo_metricas,
    bg=COLOR_SUAVE
)

recall_frame.grid(
    row=0,
    column=1,
    sticky="ew",
    padx=(6, 0)
)

tk.Label(
    recall_frame,
    text="Recall",
    bg=COLOR_SUAVE,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9, "bold")
).pack(
    anchor="w",
    padx=16,
    pady=(12, 2)
)

valor_recall = tk.Label(
    recall_frame,
    text="—",
    bg=COLOR_SUAVE,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 17, "bold")
)

valor_recall.pack(
    anchor="w",
    padx=16,
    pady=(0, 12)
)

lbl_train_test = tk.Label(
    modelo_inner,
    text="Ejecuta el análisis para ver los resultados.",
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9)
)

lbl_train_test.pack(
    anchor="w",
    pady=(0, 4)
)

lbl_variables = tk.Label(
    modelo_inner,
    text="Variables del modelo: —",
    bg=COLOR_PANEL,
    fg=COLOR_SECUNDARIO,
    font=("Segoe UI", 9)
)

lbl_variables.pack(
    anchor="w",
    pady=(0, 16)
)

tk.Label(
    modelo_inner,
    text="Matriz de confusión",
    bg=COLOR_PANEL,
    fg=COLOR_TEXTO,
    font=("Segoe UI", 11, "bold")
).pack(
    anchor="w",
    pady=(0, 8)
)

tabla_matriz = ttk.Treeview(
    modelo_inner,
    columns=("real", "pred0", "pred1"),
    show="headings",
    height=2
)

tabla_matriz.heading(
    "real",
    text=""
)

tabla_matriz.heading(
    "pred0",
    text="Predicción 0"
)

tabla_matriz.heading(
    "pred1",
    text="Predicción 1"
)

tabla_matriz.column(
    "real",
    anchor="center",
    width=180
)

tabla_matriz.column(
    "pred0",
    anchor="center",
    width=180
)

tabla_matriz.column(
    "pred1",
    anchor="center",
    width=180
)

tabla_matriz.pack(
    fill="x"
)

tk.Label(
    modelo_inner,
    text=(
        "La muestra del modelo se obtiene con Reservoir Sampling: "
        "se recorre todo el dataset sin almacenar todas sus filas en RAM."
    ),
    bg=COLOR_PANEL,
    fg="#94A3B8",
    font=("Segoe UI", 8),
    wraplength=780,
    justify="left"
).pack(
    anchor="w",
    pady=(14, 0)
)


# Estado inicial
limpiar_resultados()

ventana.mainloop()
